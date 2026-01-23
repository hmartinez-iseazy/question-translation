import pytest
from fastapi.testclient import TestClient
from unittest.mock import patch, MagicMock
import io

from src.main import app
from src.config.settings import Settings


@pytest.fixture
def mock_settings():
    """Mock settings for testing."""
    return Settings(
        deepl_api_key="test-api-key",
        api_key="test-auth-key",
        debug=True,
        log_format="text",
    )


@pytest.fixture
def client(mock_settings):
    """Test client with mocked settings."""
    with patch("src.config.settings.get_settings", return_value=mock_settings):
        with patch("src.middleware.auth.get_settings", return_value=mock_settings):
            with patch("src.middleware.rate_limit.get_settings", return_value=mock_settings):
                yield TestClient(app)


@pytest.fixture
def auth_headers(mock_settings):
    """Headers with valid API key."""
    return {"X-API-Key": mock_settings.api_key}


@pytest.fixture
def sample_excel_bytes():
    """Create a minimal valid Excel file for testing."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"

    # Row 3: language cell
    ws["C3"] = "es"

    # Row 6: headers
    ws["A6"] = "N"
    ws["D6"] = "Question"
    ws["E6"] = "Correct Answer"
    ws["F6"] = "Option 1"
    ws["G6"] = "Option 2"
    ws["H6"] = "Option 3"

    # Row 7: data
    ws["A7"] = 1
    ws["D7"] = "¿Pregunta de prueba?"
    ws["E7"] = "Respuesta correcta"
    ws["F7"] = "Opción 1"
    ws["G7"] = "Opción 2"
    ws["H7"] = "Opción 3"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


@pytest.fixture
def mock_translator():
    """Mock translator service."""
    with patch("src.services.translator.get_translator_service") as mock:
        translator = MagicMock()
        translator.check_connection.return_value = True
        translator.get_usage.return_value = {
            "character_count": 1000,
            "character_limit": 500000,
        }
        mock.return_value = translator
        yield translator
