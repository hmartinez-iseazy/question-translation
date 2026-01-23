#!/usr/bin/env python3
"""
Test script for the Question Translation API.

Usage:
    python scripts/test_api.py [excel_file] [target_language]

Examples:
    python scripts/test_api.py                           # Uses sample file, translates to French
    python scripts/test_api.py questions.xlsx EN-US      # Translate to American English
    python scripts/test_api.py questions.xlsx DE         # Translate to German
"""

import sys
import httpx
from pathlib import Path

BASE_URL = "http://localhost:8000"


def test_health():
    """Test health endpoint."""
    print("=" * 50)
    print("Testing: GET /")
    response = httpx.get(f"{BASE_URL}/")
    print(f"Status: {response.status_code}")
    print(f"Response: {response.json()}")
    return response.status_code == 200


def test_languages():
    """Test supported languages endpoint."""
    print("\n" + "=" * 50)
    print("Testing: GET /languages")
    response = httpx.get(f"{BASE_URL}/languages")
    print(f"Status: {response.status_code}")
    data = response.json()
    print(f"Total languages supported: {data['total']}")
    print("First 5 languages:")
    for lang in data["languages"][:5]:
        print(f"  - {lang['code']}: {lang['name']}")
    return response.status_code == 200


def test_usage():
    """Test API usage endpoint."""
    print("\n" + "=" * 50)
    print("Testing: GET /usage")
    response = httpx.get(f"{BASE_URL}/usage")
    print(f"Status: {response.status_code}")
    if response.status_code == 200:
        data = response.json()
        print(f"Characters used: {data['usage']['character_count']}")
        print(f"Characters limit: {data['usage']['character_limit']}")
        print(f"Remaining: {data['remaining_characters']}")
    else:
        print(f"Error: {response.json()}")
    return response.status_code == 200


def test_translate(file_path: str, target_language: str = "FR"):
    """Test translation endpoint."""
    print("\n" + "=" * 50)
    print(f"Testing: POST /translate")
    print(f"File: {file_path}")
    print(f"Target language: {target_language}")

    path = Path(file_path)
    if not path.exists():
        print(f"ERROR: File not found: {file_path}")
        return False

    with open(path, "rb") as f:
        files = {"file": (path.name, f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        data = {"target_language": target_language}

        response = httpx.post(
            f"{BASE_URL}/translate",
            files=files,
            data=data,
            timeout=120.0  # Translation can take time
        )

    print(f"Status: {response.status_code}")

    if response.status_code == 200:
        # Save translated file
        output_name = f"{path.stem}_{target_language.lower()}.xlsx"
        output_path = path.parent / output_name
        with open(output_path, "wb") as f:
            f.write(response.content)
        print(f"Translated file saved to: {output_path}")
        print(f"Source language: {response.headers.get('X-Source-Language', 'N/A')}")
        print(f"Rows translated: {response.headers.get('X-Rows-Translated', 'N/A')}")
        return True
    else:
        print(f"Error: {response.json()}")
        return False


def test_translate_info(file_path: str, target_language: str = "FR"):
    """Test translation info endpoint (no file download)."""
    print("\n" + "=" * 50)
    print(f"Testing: POST /translate/info")
    print(f"File: {file_path}")
    print(f"Target language: {target_language}")

    path = Path(file_path)
    if not path.exists():
        print(f"ERROR: File not found: {file_path}")
        return False

    with open(path, "rb") as f:
        files = {"file": (path.name, f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        data = {"target_language": target_language}

        response = httpx.post(
            f"{BASE_URL}/translate/info",
            files=files,
            data=data,
            timeout=120.0
        )

    print(f"Status: {response.status_code}")
    print(f"Response: {response.json()}")
    return response.status_code == 200


def create_sample_excel():
    """Create a sample Excel file for testing."""
    try:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Questions_process"

        # Headers
        headers = [
            "Id", "Question", "Correct Answer", "Answer 2", "Answer 3",
            "Answer 4", "Category code", "Lang", "Question Type", "Start Date"
        ]
        for col, header in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=header)

        # Sample data (questions in Spanish)
        sample_data = [
            [1, "¿Cuál es la función principal del corazón?", "Bombear sangre al cuerpo",
             "Producir hormonas", "Filtrar toxinas", "Almacenar oxígeno",
             "CAR001", "es", "single", "2024-01-01"],
            [2, "¿Qué vitamina se produce con la exposición al sol?", "Vitamina D",
             "Vitamina A", "Vitamina C", "Vitamina B12",
             "NUT001", "es", "single", "2024-01-01"],
            [3, "¿Cuántos huesos tiene el cuerpo humano adulto?", "206 huesos",
             "180 huesos", "250 huesos", "300 huesos",
             "ANA001", "es", "single", "2024-01-01"],
        ]

        for row_idx, row_data in enumerate(sample_data, 2):
            for col_idx, value in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)

        # Save
        sample_path = Path("sample_questions.xlsx")
        wb.save(sample_path)
        print(f"Sample Excel file created: {sample_path}")
        return str(sample_path)

    except ImportError:
        print("openpyxl not installed. Run: pip install openpyxl")
        return None


def main():
    print("Question Translation API Test Script")
    print("=" * 50)

    # Parse arguments
    file_path = sys.argv[1] if len(sys.argv) > 1 else None
    target_language = sys.argv[2] if len(sys.argv) > 2 else "FR"

    # Run basic tests
    all_passed = True
    all_passed &= test_health()
    all_passed &= test_languages()
    all_passed &= test_usage()

    # Create sample file if none provided
    if not file_path:
        print("\n" + "=" * 50)
        print("No file provided. Creating sample Excel file...")
        file_path = create_sample_excel()

    # Test translation if we have a file
    if file_path:
        all_passed &= test_translate_info(file_path, target_language)
        all_passed &= test_translate(file_path, target_language)

    print("\n" + "=" * 50)
    print(f"All tests passed: {all_passed}")
    return 0 if all_passed else 1


if __name__ == "__main__":
    sys.exit(main())
